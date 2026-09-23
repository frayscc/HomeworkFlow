import importlib
from urllib.parse import unquote

from fastapi.testclient import TestClient
from openpyxl import load_workbook


def test_end_to_end_issue_scan_confirm_and_export(tmp_path, monkeypatch):
    monkeypatch.setenv("HOMEWORKFLOW_DATA_DIR", str(tmp_path / "var"))
    main = importlib.import_module("homeworkflow.main")
    main = importlib.reload(main)
    client = TestClient(main.app)
    roster = "01 学生甲\n03\t学生乙\n"

    generated = client.post("/api/forms/generate", data={
        "roster_text": roster, "class_name": "演示班", "start_date": "2026-09-28", "end_date": "2026-09-28",
        "subjects": '["数学"]', "weekdays_only": "true", "copies": "1",
    })
    assert generated.status_code == 200
    manifest_name = unquote(generated.headers["x-homeworkflow-manifest"])
    assert client.get(f"/api/manifests/{manifest_name}").status_code == 200

    scanned = client.post("/api/scans/analyze",
                          files={"scan_file": ("returned.pdf", generated.content, "application/pdf")})
    assert scanned.status_code == 200, scanned.text
    assert scanned.json()["successful_count"] == 2

    week = client.get("/api/weeks").json()["weeks"][0]
    assert week["status"] == "scanned"
    confirmed = client.post(f"/api/weeks/{week['batch_id']}/confirm")
    assert confirmed.status_code == 200

    exported = client.post(f"/api/weeks/{week['batch_id']}/export")
    assert exported.status_code == 200
    path = tmp_path / "export.xlsx"
    path.write_bytes(exported.content)
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["周汇总", "缺交明细", "平台导入"]
