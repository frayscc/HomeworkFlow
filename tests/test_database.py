from datetime import date

from homeworkflow.database import Database
from homeworkflow.forms import generate_weekly_packet
from homeworkflow.roster import Roster, Student
from homeworkflow.reporting import export_week_xlsx
from openpyxl import load_workbook


def test_review_and_confirm_workflow(tmp_path):
    roster = Roster("数据库测试班", (Student("1", "学生甲"), Student("3", "学生乙")))
    pdf = tmp_path / "week.pdf"
    manifest = tmp_path / "week.manifest.json"
    bundle = generate_weekly_packet(pdf, manifest, roster=roster,
                                    start=date(2026, 9, 28), end=date(2026, 9, 28),
                                    subjects=("数学",))
    db = Database(tmp_path / "data.sqlite3")
    db.register_bundle(bundle, pdf, manifest)
    form = bundle["forms"][0]
    result = {
        "batch_id": bundle["batch_id"], "form_id": form["form_id"],
        "observations": [
            {"slot_id": "r00-d0", "student_number": "1", "student_name": "学生甲",
             "date": "2026-09-28", "classification": "blank", "confidence": 0.96,
             "features": {"ink_density": 0.0}},
            {"slot_id": "r01-d0", "student_number": "3", "student_name": "学生乙",
             "date": "2026-09-28", "classification": "review", "confidence": 0.0,
             "features": {"ink_density": 0.7}},
        ],
    }
    source = tmp_path / "scan.jpg"
    source.write_bytes(b"test")
    db.record_import("import-1", source.name, source, [result])

    weeks = db.list_weeks()
    assert weeks[0]["status"] == "scanned"
    assert weeks[0]["pending_reviews"] == 1
    review = db.pending_reviews(bundle["batch_id"])[0]
    db.review(review["id"], "missing")
    confirmed = db.confirm_week(bundle["batch_id"])
    assert confirmed["status"] == "confirmed"

    export_path = tmp_path / "report.xlsx"
    export_week_xlsx(export_path, db.report_data(bundle["batch_id"]))
    workbook = load_workbook(export_path)
    assert workbook.sheetnames == ["周汇总", "缺交明细", "平台导入"]
    assert workbook["缺交明细"].max_row == 3
    assert workbook["周汇总"].cell(2, workbook["周汇总"].max_column).value == 1


def test_platform_template_controls_columns(tmp_path):
    from openpyxl import Workbook
    from homeworkflow.reporting import build_workbook

    template = Workbook()
    template.active.append(["学生姓名", "作业日期", "科目", "自定义列"])
    stream = __import__("io").BytesIO()
    template.save(stream)
    data = {
        "week": {"class_name": "演示班"}, "students": [{"number": "1", "name": "学生甲"}],
        "subjects": ["数学"], "dates": ["2026-09-28"], "notes": [],
        "results": [{"date": "2026-09-28", "subject": "数学", "student_number": "1",
                     "student_name": "学生甲", "final_status": "missing", "raw_classifications": "blank"}],
    }
    workbook = build_workbook(data, stream.getvalue())
    assert [cell.value for cell in workbook["平台导入"][1]] == ["学生姓名", "作业日期", "科目", "自定义列"]
    assert [cell.value for cell in workbook["平台导入"][2]] == ["学生甲", "2026-09-28", "数学", ""]
