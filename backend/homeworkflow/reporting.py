from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


HEADER_FILL = PatternFill("solid", fgColor="2457D6")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ALIASES = {
    "日期": "date", "作业日期": "date", "学科": "subject", "科目": "subject",
    "学号": "student_number", "编号": "student_number", "姓名": "student_name",
    "学生姓名": "student_name", "状态": "status_cn", "提交状态": "status_cn",
    "班级": "class_name", "备注": "note", "原始识别": "raw_classifications",
}


def _style_sheet(sheet, widths: dict[str, float] | None = None) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    if widths:
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width


def _missing_rows(data: dict[str, Any]) -> list[dict[str, Any]]:
    class_name = data["week"]["class_name"]
    return [{**item, "class_name": class_name, "status_cn": "缺交", "note": ""}
            for item in data["results"] if item["final_status"] == "missing"]


def build_workbook(data: dict[str, Any], template_content: bytes | None = None) -> Workbook:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "周汇总"
    summary.append(["学号", "姓名", *[f"{subject}缺交" for subject in data["subjects"]], "缺交总次数"])
    missing = _missing_rows(data)
    for student in data["students"]:
        counts = [sum(item["student_number"] == student["number"] and item["subject"] == subject
                      for item in missing) for subject in data["subjects"]]
        summary.append([student["number"], student["name"], *counts, sum(counts)])
    _style_sheet(summary, {"A": 10, "B": 16})

    detail = workbook.create_sheet("缺交明细")
    detail.append(["日期", "学科", "学号", "姓名", "状态", "原始识别", "备注"])
    for item in sorted(missing, key=lambda row: (row["date"], row["subject"], row["student_number"])):
        detail.append([item["date"], item["subject"], item["student_number"], item["student_name"],
                       item["status_cn"], item["raw_classifications"], item["note"]])
    _style_sheet(detail, {"A": 13, "B": 10, "C": 10, "D": 16, "E": 10, "F": 18, "G": 28})

    if template_content:
        template = load_workbook(io.BytesIO(template_content))
        source = template.active
        headers = [cell.value for cell in source[1] if cell.value is not None]
        if not headers:
            raise ValueError("平台模板第一行没有列名")
    else:
        headers = ["日期", "学科", "学号", "姓名", "状态", "备注"]
    platform = workbook.create_sheet("平台导入")
    platform.append(headers)
    for item in sorted(missing, key=lambda row: (row["date"], row["subject"], row["student_number"])):
        platform.append([item.get(ALIASES.get(str(header), ""), "") for header in headers])
    _style_sheet(platform)
    for column in platform.columns:
        letter = column[0].column_letter
        platform.column_dimensions[letter].width = min(32, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
    return workbook


def export_week_xlsx(path: Path, data: dict[str, Any], template_content: bytes | None = None) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(data, template_content).save(path)
    return path

